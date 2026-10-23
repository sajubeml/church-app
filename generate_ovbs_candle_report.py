import json
import os
import shutil
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

# 1. Load Backup JSON
BACKUP_PATH = r"C:\Users\sajub\Downloads\St_Gregorios_Church_Backup_2026-10-23.json"
with open(BACKUP_PATH, 'r', encoding='utf-8') as f:
    data = json.load(f)

cb = data.get('cashbook', [])

ovbs_list = []
candle_cup_list = []
candles_other_list = []

for idx, r in enumerate(cb):
    rem = str(r.get('G') or '').strip()
    rem_lower = rem.lower()
    head = str(r.get('E') or '').strip()
    code = str(r.get('F') or '').strip()
    dt = str(r.get('A') or '').strip()
    rcpt = str(r.get('B') or '').strip()
    reg = str(r.get('C') or '').strip()
    name = str(r.get('D') or '').strip()
    cash = float(r.get('H') or 0) if str(r.get('H') or '').strip() else 0.0
    bank = float(r.get('I') or 0) if str(r.get('I') or '').strip() else 0.0
    amt = cash + bank
    
    row_data = {
        'row': idx, 'date': dt, 'rcpt': rcpt, 'reg': reg, 'name': name,
        'rem': rem, 'head': head, 'code': code, 'cash': cash, 'bank': bank, 'amt': amt
    }
    
    if code == 'RP-3.66' or 'ovbs' in head.lower():
        ovbs_list.append(row_data)
    elif 'candle cup' in rem_lower:
        candle_cup_list.append(row_data)
    elif any(k in rem_lower for k in ['candle', 'candles']) and 'sale' not in rem_lower and name != '':
        candles_other_list.append(row_data)

# Sort by date / receipt
def sort_key(x):
    return (x['date'], int(x['rcpt']) if x['rcpt'].isdigit() else 999999)

ovbs_list.sort(key=sort_key)
candle_cup_list.sort(key=sort_key)
candles_other_list.sort(key=sort_key)

tot_ovbs_cash = sum(x['cash'] for x in ovbs_list)
tot_ovbs_bank = sum(x['bank'] for x in ovbs_list)
tot_ovbs_amt = sum(x['amt'] for x in ovbs_list)

tot_cc_cash = sum(x['cash'] for x in candle_cup_list)
tot_cc_bank = sum(x['bank'] for x in candle_cup_list)
tot_cc_amt = sum(x['amt'] for x in candle_cup_list)

tot_oc_cash = sum(x['cash'] for x in candles_other_list)
tot_oc_bank = sum(x['bank'] for x in candles_other_list)
tot_oc_amt = sum(x['amt'] for x in candles_other_list)

candle_sub_cash = tot_cc_cash + tot_oc_cash
candle_sub_bank = tot_cc_bank + tot_oc_bank
candle_sub_amt = tot_cc_amt + tot_oc_amt

grand_cash = tot_ovbs_cash + tot_cc_cash + tot_oc_cash
grand_bank = tot_ovbs_bank + tot_cc_bank + tot_oc_bank
grand_amt = tot_ovbs_amt + tot_cc_amt + tot_oc_amt
grand_count = len(ovbs_list) + len(candle_cup_list) + len(candles_other_list)

# ==========================================
# PART 1: GENERATE EXCEL WORKBOOK (.XLSX)
# ==========================================
wb = openpyxl.Workbook()
wb.remove(wb.active)

navy_dark = "1B365D"
navy_light = "2E5B88"
border_color = "CBD5E1"

font_title = Font(name="Calibri", size=16, bold=True, color="1B365D")
font_sub = Font(name="Calibri", size=11, italic=True, color="475569")
font_sec = Font(name="Calibri", size=12, bold=True, color="1B365D")
font_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="Calibri", size=10, bold=True)
font_regular = Font(name="Calibri", size=10)

fill_hdr = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
fill_subhdr = PatternFill(start_color=navy_light, end_color=navy_light, fill_type="solid")
fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_tot = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

thin_border_side = Side(style='thin', color=border_color)
thick_top_side = Side(style='thin', color="1E293B")
double_bottom_side = Side(style='double', color="1E293B")

cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
total_border = Border(left=thin_border_side, right=thin_border_side, top=thick_top_side, bottom=double_bottom_side)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# --- SHEET 1: EXECUTIVE SUMMARY ---
ws_sum = wb.create_sheet(title="Executive Summary")
ws_sum.views.sheetView[0].showGridLines = True

ws_sum.merge_cells("A1:G1")
ws_sum["A1"] = "ST. GREGORIOS ORTHODOX SYRIAN CHURCH"
ws_sum["A1"].font = font_title
ws_sum["A1"].alignment = align_center

ws_sum.merge_cells("A2:G2")
ws_sum["A2"] = "OVBS & CANDLE CUPS / CANDLE INDIVIDUAL CONTRIBUTION REPORT"
ws_sum["A2"].font = Font(name="Calibri", size=12, bold=True, color="334155")
ws_sum["A2"].alignment = align_center

ws_sum.merge_cells("A3:G3")
ws_sum["A3"] = f"Period: Past 2 Months (August 2026 - October 2026) | Generated: {datetime.now().strftime('%d-%b-%Y')}"
ws_sum["A3"].font = font_sub
ws_sum["A3"].alignment = align_center

headers_sum = ["Category / Account Head", "Accounting Code", "Contributors Count", "Cash (Rs.)", "Bank / Online (Rs.)", "Total Amount (Rs.)", "% Share"]
row_curr = 5
for col_idx, h in enumerate(headers_sum, 1):
    c = ws_sum.cell(row=row_curr, column=col_idx, value=h)
    c.font = font_hdr
    c.fill = fill_hdr
    c.alignment = align_center if col_idx in [2, 3] else (align_right if col_idx >= 4 else align_left)
    c.border = cell_border

sum_rows = [
    ("1. OVBS (Orthodox Vacation Bible School)", "RP-3.66", len(ovbs_list), tot_ovbs_cash, tot_ovbs_bank, tot_ovbs_amt, tot_ovbs_amt / grand_amt),
    ("2. Candle Cup Sponsorships", "RP-2.02", len(candle_cup_list), tot_cc_cash, tot_cc_bank, tot_cc_amt, tot_cc_amt / grand_amt),
    ("3. Other Candle Donations / Sponsorships", "RP-2.02", len(candles_other_list), tot_oc_cash, tot_oc_bank, tot_oc_amt, tot_oc_amt / grand_amt),
]

for r_data in sum_rows:
    row_curr += 1
    for c_idx, val in enumerate(r_data, 1):
        c = ws_sum.cell(row=row_curr, column=c_idx, value=val)
        c.font = font_bold if c_idx == 1 else font_regular
        c.border = cell_border
        if c_idx == 1: c.alignment = align_left
        elif c_idx in [2, 3]: c.alignment = align_center
        elif c_idx in [4, 5, 6]:
            c.alignment = align_right
            c.number_format = "#,##0.00"
        elif c_idx == 7:
            c.alignment = align_right
            c.number_format = "0.0%"

row_curr += 1
sub_row = ("   ↳ Subtotal: All Candle & Candle Cup Contributions", "RP-2.02", len(candle_cup_list) + len(candles_other_list), candle_sub_cash, candle_sub_bank, candle_sub_amt, candle_sub_amt / grand_amt)
for c_idx, val in enumerate(sub_row, 1):
    c = ws_sum.cell(row=row_curr, column=c_idx, value=val)
    c.font = font_bold
    c.fill = fill_alt
    c.border = cell_border
    if c_idx == 1: c.alignment = align_left
    elif c_idx in [2, 3]: c.alignment = align_center
    elif c_idx in [4, 5, 6]:
        c.alignment = align_right
        c.number_format = "#,##0.00"
    elif c_idx == 7:
        c.alignment = align_right
        c.number_format = "0.0%"

row_curr += 1
tot_row = ("GRAND TOTAL CONTRIBUTIONS", "ALL", grand_count, grand_cash, grand_bank, grand_amt, 1.0)
for c_idx, val in enumerate(tot_row, 1):
    c = ws_sum.cell(row=row_curr, column=c_idx, value=val)
    c.font = font_bold
    c.fill = fill_tot
    c.border = total_border
    if c_idx == 1: c.alignment = align_left
    elif c_idx in [2, 3]: c.alignment = align_center
    elif c_idx in [4, 5, 6]:
        c.alignment = align_right
        c.number_format = "#,##0.00"
    elif c_idx == 7:
        c.alignment = align_right
        c.number_format = "0.0%"

row_curr += 3
ws_sum.cell(row=row_curr, column=1, value="KEY TAKEAWAYS:").font = font_sec
row_curr += 1
ws_sum.cell(row=row_curr, column=1, value="• OVBS had 24 contributions totaling Rs. 62,000.00 (Cash: Rs. 55,000.00, Bank: Rs. 7,000.00).").font = font_regular
row_curr += 1
ws_sum.cell(row=row_curr, column=1, value="• Candle Cup sponsorships had 7 distinct entries totaling Rs. 9,500.00 (Cash: Rs. 3,000.00, Bank: Rs. 6,500.00).").font = font_regular
row_curr += 1
ws_sum.cell(row=row_curr, column=1, value="• Other Candle individual donations (04-10-2026) had 6 entries totaling Rs. 8,500.00 (Cash: Rs. 6,500.00, Bank: Rs. 2,000.00).").font = font_regular
row_curr += 1
ws_sum.cell(row=row_curr, column=1, value="• Combined total for all 37 contributions is Rs. 80,000.00 (Cash: Rs. 64,500.00 [80.6%], Bank: Rs. 15,500.00 [19.4%]).").font = font_regular


# Helper to build tabular contribution sheets
def create_contribution_sheet(sheet_title, report_title, item_list, code_text):
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells("A1:I1")
    ws["A1"] = "ST. GREGORIOS ORTHODOX SYRIAN CHURCH"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:I2")
    ws["A2"] = report_title
    ws["A2"].font = Font(name="Calibri", size=12, bold=True, color="334155")
    ws["A2"].alignment = align_center

    ws.merge_cells("A3:I3")
    ws["A3"] = f"Account Head Code: {code_text} | Total Records: {len(item_list)} | Generated: {datetime.now().strftime('%d-%b-%Y')}"
    ws["A3"].font = font_sub
    ws["A3"].alignment = align_center
    
    cols = ["Sl No.", "Date", "Receipt No.", "Reg No.", "Member / Contributor Name", "Remarks / Dedication", "Cash (Rs.)", "Bank (Rs.)", "Total (Rs.)"]
    start_row = 5
    for c_idx, h in enumerate(cols, 1):
        c = ws.cell(row=start_row, column=c_idx, value=h)
        c.font = font_hdr
        c.fill = fill_hdr
        c.border = cell_border
        c.alignment = align_center if c_idx in [1, 2, 3, 4] else (align_right if c_idx >= 7 else align_left)
    
    curr = start_row
    for idx, item in enumerate(item_list, 1):
        curr += 1
        is_even = (idx % 2 == 0)
        row_vals = [
            idx,
            item['date'],
            item['rcpt'],
            item['reg'] if item['reg'] else "-",
            item['name'],
            item['rem'] if item['rem'] else "-",
            item['cash'],
            item['bank'],
            item['amt']
        ]
        for c_idx, val in enumerate(row_vals, 1):
            c = ws.cell(row=curr, column=c_idx, value=val)
            c.font = font_regular
            c.border = cell_border
            if is_even: c.fill = fill_alt
            if c_idx in [1, 2, 3, 4]: c.alignment = align_center
            elif c_idx in [5, 6]: c.alignment = align_left
            elif c_idx in [7, 8, 9]:
                c.alignment = align_right
                c.number_format = "#,##0.00"
    
    curr += 1
    total_cash_formula = f"=SUM(G{start_row+1}:G{curr-1})"
    total_bank_formula = f"=SUM(H{start_row+1}:H{curr-1})"
    total_amt_formula = f"=SUM(I{start_row+1}:I{curr-1})"
    
    ws.merge_cells(f"A{curr}:F{curr}")
    tot_label = ws.cell(row=curr, column=1, value="TOTAL")
    tot_label.font = font_bold
    tot_label.alignment = align_right
    for c_idx in range(1, 7):
        ws.cell(row=curr, column=c_idx).border = total_border
        ws.cell(row=curr, column=c_idx).fill = fill_tot
        
    c_cash = ws.cell(row=curr, column=7, value=total_cash_formula)
    c_cash.font = font_bold; c_cash.alignment = align_right; c_cash.number_format = "#,##0.00"; c_cash.border = total_border; c_cash.fill = fill_tot

    c_bank = ws.cell(row=curr, column=8, value=total_bank_formula)
    c_bank.font = font_bold; c_bank.alignment = align_right; c_bank.number_format = "#,##0.00"; c_bank.border = total_border; c_bank.fill = fill_tot

    c_amt = ws.cell(row=curr, column=9, value=total_amt_formula)
    c_amt.font = font_bold; c_amt.alignment = align_right; c_amt.number_format = "#,##0.00"; c_amt.border = total_border; c_amt.fill = fill_tot

# --- SHEET 2: OVBS CONTRIBUTIONS ---
create_contribution_sheet("OVBS Contributions", "OVBS INDIVIDUAL CONTRIBUTIONS REPORT", ovbs_list, "RP-3.66 (OVBS)")

# --- SHEET 3: CANDLE CUPS & CANDLES ---
ws_cc = wb.create_sheet(title="Candle Cups & Candles")
ws_cc.views.sheetView[0].showGridLines = True

ws_cc.merge_cells("A1:I1")
ws_cc["A1"] = "ST. GREGORIOS ORTHODOX SYRIAN CHURCH"
ws_cc["A1"].font = font_title
ws_cc["A1"].alignment = align_center

ws_cc.merge_cells("A2:I2")
ws_cc["A2"] = "CANDLE CUPS & CANDLE SPONSORSHIPS REPORT"
ws_cc["A2"].font = Font(name="Calibri", size=12, bold=True, color="334155")
ws_cc["A2"].alignment = align_center

ws_cc.merge_cells("A3:I3")
ws_cc["A3"] = f"Account Head: Donation General (Code: RP-2.02) | Generated: {datetime.now().strftime('%d-%b-%Y')}"
ws_cc["A3"].font = font_sub
ws_cc["A3"].alignment = align_center

cols = ["Sl No.", "Date", "Receipt No.", "Reg No.", "Member / Contributor Name", "Remarks / Dedication", "Cash (Rs.)", "Bank (Rs.)", "Total (Rs.)"]

# SECTION 1: CANDLE CUP SPONSORSHIPS
row_c = 5
ws_cc.cell(row=row_c, column=1, value="SECTION A: CANDLE CUP SPONSORSHIPS (Remarks: 'Candle Cup')").font = font_sec
row_c += 1
for c_idx, h in enumerate(cols, 1):
    c = ws_cc.cell(row=row_c, column=c_idx, value=h)
    c.font = font_hdr
    c.fill = fill_subhdr
    c.border = cell_border
    c.alignment = align_center if c_idx in [1, 2, 3, 4] else (align_right if c_idx >= 7 else align_left)

sec_a_start = row_c + 1
for idx, item in enumerate(candle_cup_list, 1):
    row_c += 1
    is_even = (idx % 2 == 0)
    row_vals = [idx, item['date'], item['rcpt'], item['reg'] or "-", item['name'], item['rem'] or "-", item['cash'], item['bank'], item['amt']]
    for c_idx, val in enumerate(row_vals, 1):
        c = ws_cc.cell(row=row_c, column=c_idx, value=val)
        c.font = font_regular
        c.border = cell_border
        if is_even: c.fill = fill_alt
        if c_idx in [1, 2, 3, 4]: c.alignment = align_center
        elif c_idx in [5, 6]: c.alignment = align_left
        elif c_idx in [7, 8, 9]:
            c.alignment = align_right
            c.number_format = "#,##0.00"

row_c += 1
ws_cc.merge_cells(f"A{row_c}:F{row_c}")
sub_a_label = ws_cc.cell(row=row_c, column=1, value="SUBTOTAL - CANDLE CUPS")
sub_a_label.font = font_bold
sub_a_label.alignment = align_right
for c_idx in range(1, 7):
    ws_cc.cell(row=row_c, column=c_idx).border = total_border
    ws_cc.cell(row=row_c, column=c_idx).fill = fill_alt

c_a_cash = ws_cc.cell(row=row_c, column=7, value=f"=SUM(G{sec_a_start}:G{row_c-1})")
c_a_cash.font = font_bold; c_a_cash.alignment = align_right; c_a_cash.number_format = "#,##0.00"; c_a_cash.border = total_border; c_a_cash.fill = fill_alt
c_a_bank = ws_cc.cell(row=row_c, column=8, value=f"=SUM(H{sec_a_start}:H{row_c-1})")
c_a_bank.font = font_bold; c_a_bank.alignment = align_right; c_a_bank.number_format = "#,##0.00"; c_a_bank.border = total_border; c_a_bank.fill = fill_alt
c_a_amt = ws_cc.cell(row=row_c, column=9, value=f"=SUM(I{sec_a_start}:I{row_c-1})")
c_a_amt.font = font_bold; c_a_amt.alignment = align_right; c_a_amt.number_format = "#,##0.00"; c_a_amt.border = total_border; c_a_amt.fill = fill_alt
sec_a_tot_row = row_c

# SECTION 2: OTHER INDIVIDUAL CANDLE DONATIONS
row_c += 3
ws_cc.cell(row=row_c, column=1, value="SECTION B: ADDITIONAL CANDLE SPONSORSHIPS (Remarks: 'Candles')").font = font_sec
row_c += 1
for c_idx, h in enumerate(cols, 1):
    c = ws_cc.cell(row=row_c, column=c_idx, value=h)
    c.font = font_hdr
    c.fill = fill_subhdr
    c.border = cell_border
    c.alignment = align_center if c_idx in [1, 2, 3, 4] else (align_right if c_idx >= 7 else align_left)

sec_b_start = row_c + 1
for idx, item in enumerate(candles_other_list, 1):
    row_c += 1
    is_even = (idx % 2 == 0)
    row_vals = [idx, item['date'], item['rcpt'], item['reg'] or "-", item['name'], item['rem'] or "-", item['cash'], item['bank'], item['amt']]
    for c_idx, val in enumerate(row_vals, 1):
        c = ws_cc.cell(row=row_c, column=c_idx, value=val)
        c.font = font_regular
        c.border = cell_border
        if is_even: c.fill = fill_alt
        if c_idx in [1, 2, 3, 4]: c.alignment = align_center
        elif c_idx in [5, 6]: c.alignment = align_left
        elif c_idx in [7, 8, 9]:
            c.alignment = align_right
            c.number_format = "#,##0.00"

row_c += 1
ws_cc.merge_cells(f"A{row_c}:F{row_c}")
sub_b_label = ws_cc.cell(row=row_c, column=1, value="SUBTOTAL - OTHER CANDLES")
sub_b_label.font = font_bold
sub_b_label.alignment = align_right
for c_idx in range(1, 7):
    ws_cc.cell(row=row_c, column=c_idx).border = total_border
    ws_cc.cell(row=row_c, column=c_idx).fill = fill_alt

c_b_cash = ws_cc.cell(row=row_c, column=7, value=f"=SUM(G{sec_b_start}:G{row_c-1})")
c_b_cash.font = font_bold; c_b_cash.alignment = align_right; c_b_cash.number_format = "#,##0.00"; c_b_cash.border = total_border; c_b_cash.fill = fill_alt
c_b_bank = ws_cc.cell(row=row_c, column=8, value=f"=SUM(H{sec_b_start}:H{row_c-1})")
c_b_bank.font = font_bold; c_b_bank.alignment = align_right; c_b_bank.number_format = "#,##0.00"; c_b_bank.border = total_border; c_b_bank.fill = fill_alt
c_b_amt = ws_cc.cell(row=row_c, column=9, value=f"=SUM(I{sec_b_start}:I{row_c-1})")
c_b_amt.font = font_bold; c_b_amt.alignment = align_right; c_b_amt.number_format = "#,##0.00"; c_b_amt.border = total_border; c_b_amt.fill = fill_alt
sec_b_tot_row = row_c

# GRAND TOTAL COMBINED CANDLES
row_c += 2
ws_cc.merge_cells(f"A{row_c}:F{row_c}")
tot_c_label = ws_cc.cell(row=row_c, column=1, value="TOTAL ALL CANDLE CONTRIBUTIONS (A + B)")
tot_c_label.font = font_bold
tot_c_label.alignment = align_right
for c_idx in range(1, 7):
    ws_cc.cell(row=row_c, column=c_idx).border = total_border
    ws_cc.cell(row=row_c, column=c_idx).fill = fill_tot

c_tot_cash = ws_cc.cell(row=row_c, column=7, value=f"=G{sec_a_tot_row}+G{sec_b_tot_row}")
c_tot_cash.font = font_bold; c_tot_cash.alignment = align_right; c_tot_cash.number_format = "#,##0.00"; c_tot_cash.border = total_border; c_tot_cash.fill = fill_tot
c_tot_bank = ws_cc.cell(row=row_c, column=8, value=f"=H{sec_a_tot_row}+H{sec_b_tot_row}")
c_tot_bank.font = font_bold; c_tot_bank.alignment = align_right; c_tot_bank.number_format = "#,##0.00"; c_tot_bank.border = total_border; c_tot_bank.fill = fill_tot
c_tot_amt = ws_cc.cell(row=row_c, column=9, value=f"=I{sec_a_tot_row}+I{sec_b_tot_row}")
c_tot_amt.font = font_bold; c_tot_amt.alignment = align_right; c_tot_amt.number_format = "#,##0.00"; c_tot_amt.border = total_border; c_tot_amt.fill = fill_tot


# --- SHEET 4: ALL 37 CONTRIBUTIONS (CONSOLIDATED) ---
ws_all = wb.create_sheet(title="All Combined List")
ws_all.views.sheetView[0].showGridLines = True

ws_all.merge_cells("A1:J1")
ws_all["A1"] = "ST. GREGORIOS ORTHODOX SYRIAN CHURCH"
ws_all["A1"].font = font_title
ws_all["A1"].alignment = align_center

ws_all.merge_cells("A2:J2")
ws_all["A2"] = "CONSOLIDATED INDIVIDUAL CONTRIBUTIONS: OVBS & CANDLE CUPS"
ws_all["A2"].font = Font(name="Calibri", size=12, bold=True, color="334155")
ws_all["A2"].alignment = align_center

ws_all.merge_cells("A3:J3")
ws_all["A3"] = f"Total Contributions: {grand_count} | Total Amount: Rs. {grand_amt:,.2f} | Generated: {datetime.now().strftime('%d-%b-%Y')}"
ws_all["A3"].font = font_sub
ws_all["A3"].alignment = align_center

cols_all = ["Sl No.", "Category", "Date", "Receipt No.", "Reg No.", "Member / Contributor Name", "Remarks / Dedication", "Cash (Rs.)", "Bank (Rs.)", "Total (Rs.)"]
start_row_all = 5
for c_idx, h in enumerate(cols_all, 1):
    c = ws_all.cell(row=start_row_all, column=c_idx, value=h)
    c.font = font_hdr
    c.fill = fill_hdr
    c.border = cell_border
    c.alignment = align_center if c_idx in [1, 3, 4, 5] else (align_right if c_idx >= 8 else align_left)

all_combined = []
for x in ovbs_list:
    item = dict(x); item['category'] = "OVBS"; all_combined.append(item)
for x in candle_cup_list:
    item = dict(x); item['category'] = "Candle Cup"; all_combined.append(item)
for x in candles_other_list:
    item = dict(x); item['category'] = "Candles"; all_combined.append(item)

all_combined.sort(key=sort_key)

curr_all = start_row_all
for idx, item in enumerate(all_combined, 1):
    curr_all += 1
    is_even = (idx % 2 == 0)
    row_vals = [
        idx,
        item['category'],
        item['date'],
        item['rcpt'],
        item['reg'] or "-",
        item['name'],
        item['rem'] or "-",
        item['cash'],
        item['bank'],
        item['amt']
    ]
    for c_idx, val in enumerate(row_vals, 1):
        c = ws_all.cell(row=curr_all, column=c_idx, value=val)
        c.font = font_regular
        c.border = cell_border
        if is_even: c.fill = fill_alt
        if c_idx in [1, 3, 4, 5]: c.alignment = align_center
        elif c_idx == 2:
            c.alignment = align_center
            if item['category'] == 'OVBS':
                c.font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
            elif item['category'] == 'Candle Cup':
                c.font = Font(name="Calibri", size=10, bold=True, color="92400E")
            else:
                c.font = Font(name="Calibri", size=10, bold=True, color="065F46")
        elif c_idx in [6, 7]: c.alignment = align_left
        elif c_idx in [8, 9, 10]:
            c.alignment = align_right
            c.number_format = "#,##0.00"

curr_all += 1
ws_all.merge_cells(f"A{curr_all}:G{curr_all}")
tot_all_label = ws_all.cell(row=curr_all, column=1, value="GRAND TOTAL ALL CONTRIBUTIONS")
tot_all_label.font = font_bold
tot_all_label.alignment = align_right
for c_idx in range(1, 8):
    ws_all.cell(row=curr_all, column=c_idx).border = total_border
    ws_all.cell(row=curr_all, column=c_idx).fill = fill_tot

c_all_cash = ws_all.cell(row=curr_all, column=8, value=f"=SUM(H{start_row_all+1}:H{curr_all-1})")
c_all_cash.font = font_bold; c_all_cash.alignment = align_right; c_all_cash.number_format = "#,##0.00"; c_all_cash.border = total_border; c_all_cash.fill = fill_tot

c_all_bank = ws_all.cell(row=curr_all, column=9, value=f"=SUM(I{start_row_all+1}:I{curr_all-1})")
c_all_bank.font = font_bold; c_all_bank.alignment = align_right; c_all_bank.number_format = "#,##0.00"; c_all_bank.border = total_border; c_all_bank.fill = fill_tot

c_all_amt = ws_all.cell(row=curr_all, column=10, value=f"=SUM(J{start_row_all+1}:J{curr_all-1})")
c_all_amt.font = font_bold; c_all_amt.alignment = align_right; c_all_amt.number_format = "#,##0.00"; c_all_amt.border = total_border; c_all_amt.fill = fill_tot

# Column width auto-adjust
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in [1, 2, 3]: continue
            val_str = str(cell.value or '')
            if cell.number_format and "#" in cell.number_format:
                val_str = "Rs. 99,999.00"
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

XLSX_PATH_APP = r"c:\CASHBOOK_APP\St_Gregorios_OVBS_and_Candle_Cups_Report.xlsx"
XLSX_PATH_DL = r"C:\Users\sajub\Downloads\St_Gregorios_OVBS_and_Candle_Cups_Report.xlsx"

wb.save(XLSX_PATH_APP)
shutil.copy(XLSX_PATH_APP, XLSX_PATH_DL)
print(f"Excel report saved.")


# ==========================================
# PART 2: GENERATE PDF REPORT (.PDF)
# ==========================================
PDF_PATH_APP = r"c:\CASHBOOK_APP\St_Gregorios_OVBS_and_Candle_Cups_Report.pdf"
PDF_PATH_DL = r"C:\Users\sajub\Downloads\St_Gregorios_OVBS_and_Candle_Cups_Report.pdf"

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header (page 2)
        if self._pageNumber > 1:
            self.drawString(36, 808, "St. Gregorios Orthodox Syrian Church - OVBS & Candle Cups Contribution Report")
            self.drawRightString(559, 808, f"Date: {datetime.now().strftime('%d-%b-%Y')}")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(36, 802, 559, 802)
            
        # Footer
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 38, 559, 38)
        
        self.drawString(36, 26, "Generated from Monolithic Backup: St_Gregorios_Church_Backup_2026-10-23.json")
        self.drawRightString(559, 26, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()

doc = SimpleDocTemplate(
    PDF_PATH_APP,
    pagesize=A4,
    leftMargin=36,
    rightMargin=36,
    topMargin=40,
    bottomMargin=45
)

styles = getSampleStyleSheet()

style_title = ParagraphStyle(
    'DocTitle',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=14,
    leading=17,
    alignment=1,
    textColor=colors.HexColor('#1B365D')
)

style_subtitle = ParagraphStyle(
    'DocSubTitle',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=10,
    leading=13,
    alignment=1,
    textColor=colors.HexColor('#334155')
)

style_meta = ParagraphStyle(
    'DocMeta',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=8,
    leading=10,
    alignment=1,
    textColor=colors.HexColor('#64748B')
)

style_sec_heading = ParagraphStyle(
    'SectionHeading',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=9.5,
    leading=12,
    textColor=colors.HexColor('#1B365D'),
    spaceBefore=6,
    spaceAfter=3
)

style_tbl_hdr = ParagraphStyle(
    'TblHdr',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=7,
    leading=8.5,
    alignment=1,
    textColor=colors.white
)

style_tbl_cell = ParagraphStyle(
    'TblCell',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=7,
    leading=8.5,
    textColor=colors.HexColor('#0F172A')
)

style_tbl_cell_center = ParagraphStyle('TblCellCenter', parent=style_tbl_cell, alignment=1)
style_tbl_cell_right = ParagraphStyle('TblCellRight', parent=style_tbl_cell, alignment=2)
style_tbl_cell_bold_right = ParagraphStyle('TblCellBoldRight', parent=style_tbl_cell, fontName='Helvetica-Bold', alignment=2)

story = []

# Church Header Block with Logo
logo_path = r"c:\CASHBOOK_APP\church_logo.png"
if os.path.exists(logo_path):
    logo_img = Image(logo_path, width=40, height=40)
    header_data = [
        [logo_img, [
            Paragraph("ST. GREGORIOS ORTHODOX SYRIAN CHURCH", style_title),
            Paragraph("OVBS & CANDLE CUPS INDIVIDUAL CONTRIBUTION REPORT", style_subtitle),
            Paragraph(f"Period: Past 2 Months (August 2026 - October 2026) | Generated: {datetime.now().strftime('%d-%b-%Y')}", style_meta)
        ]]
    ]
    header_tbl = Table(header_data, colWidths=[46, 477])
    header_tbl.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (0,0), 'CENTER'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(header_tbl)
else:
    story.append(Paragraph("ST. GREGORIOS ORTHODOX SYRIAN CHURCH", style_title))
    story.append(Paragraph("OVBS & CANDLE CUPS INDIVIDUAL CONTRIBUTION REPORT", style_subtitle))
    story.append(Paragraph(f"Period: Past 2 Months (August 2026 - October 2026) | Generated: {datetime.now().strftime('%d-%b-%Y')}", style_meta))

story.append(Spacer(1, 6))
story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1B365D"), spaceBefore=0, spaceAfter=6))

# EXECUTIVE SUMMARY TABLE
story.append(Paragraph("1. EXECUTIVE SUMMARY & BREAKDOWN", style_sec_heading))

summary_pdf_data = [
    [
        Paragraph("<b>Category / Head</b>", style_tbl_hdr),
        Paragraph("<b>Code</b>", style_tbl_hdr),
        Paragraph("<b>Count</b>", style_tbl_hdr),
        Paragraph("<b>Cash (Rs.)</b>", style_tbl_hdr),
        Paragraph("<b>Bank / Online (Rs.)</b>", style_tbl_hdr),
        Paragraph("<b>Total (Rs.)</b>", style_tbl_hdr),
        Paragraph("<b>Share</b>", style_tbl_hdr),
    ],
    [
        Paragraph("<b>1. OVBS Contributions</b>", style_tbl_cell),
        Paragraph("RP-3.66", style_tbl_cell_center),
        Paragraph(str(len(ovbs_list)), style_tbl_cell_center),
        Paragraph(f"{tot_ovbs_cash:,.2f}", style_tbl_cell_right),
        Paragraph(f"{tot_ovbs_bank:,.2f}", style_tbl_cell_right),
        Paragraph(f"<b>{tot_ovbs_amt:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"{(tot_ovbs_amt/grand_amt)*100:.1f}%", style_tbl_cell_center),
    ],
    [
        Paragraph("<b>2. Candle Cup Sponsorships</b>", style_tbl_cell),
        Paragraph("RP-2.02", style_tbl_cell_center),
        Paragraph(str(len(candle_cup_list)), style_tbl_cell_center),
        Paragraph(f"{tot_cc_cash:,.2f}", style_tbl_cell_right),
        Paragraph(f"{tot_cc_bank:,.2f}", style_tbl_cell_right),
        Paragraph(f"<b>{tot_cc_amt:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"{(tot_cc_amt/grand_amt)*100:.1f}%", style_tbl_cell_center),
    ],
    [
        Paragraph("<b>3. Other Candle Donations</b>", style_tbl_cell),
        Paragraph("RP-2.02", style_tbl_cell_center),
        Paragraph(str(len(candles_other_list)), style_tbl_cell_center),
        Paragraph(f"{tot_oc_cash:,.2f}", style_tbl_cell_right),
        Paragraph(f"{tot_oc_bank:,.2f}", style_tbl_cell_right),
        Paragraph(f"<b>{tot_oc_amt:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"{(tot_oc_amt/grand_amt)*100:.1f}%", style_tbl_cell_center),
    ],
    [
        Paragraph("<b>SUBTOTAL: ALL CANDLES</b>", style_tbl_cell),
        Paragraph("RP-2.02", style_tbl_cell_center),
        Paragraph(f"<b>{len(candle_cup_list)+len(candles_other_list)}</b>", style_tbl_cell_center),
        Paragraph(f"<b>{candle_sub_cash:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"<b>{candle_sub_bank:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"<b>{candle_sub_amt:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"<b>{(candle_sub_amt/grand_amt)*100:.1f}%</b>", style_tbl_cell_center),
    ],
    [
        Paragraph("<b>GRAND TOTAL CONTRIBUTIONS</b>", style_tbl_cell),
        Paragraph("ALL", style_tbl_cell_center),
        Paragraph(f"<b>{grand_count}</b>", style_tbl_cell_center),
        Paragraph(f"<b>{grand_cash:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"<b>{grand_bank:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph(f"<b>{grand_amt:,.2f}</b>", style_tbl_cell_bold_right),
        Paragraph("<b>100.0%</b>", style_tbl_cell_center),
    ]
]

summary_tbl = Table(summary_pdf_data, colWidths=[150, 48, 40, 75, 85, 80, 45])
summary_tbl.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1B365D')),
    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ('TOPPADDING', (0,0), (-1,-1), 3),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
    ('BACKGROUND', (0,4), (-1,4), colors.HexColor('#F1F5F9')),
    ('BACKGROUND', (0,5), (-1,5), colors.HexColor('#E2E8F0')),
]))
story.append(summary_tbl)
story.append(Spacer(1, 8))

# 2. OVBS CONTRIBUTIONS TABLE
story.append(Paragraph(f"2. OVBS INDIVIDUAL CONTRIBUTIONS ({len(ovbs_list)} Records | Total: Rs. {tot_ovbs_amt:,.2f})", style_sec_heading))

ovbs_pdf_data = [
    [
        Paragraph("<b>#</b>", style_tbl_hdr),
        Paragraph("<b>Date</b>", style_tbl_hdr),
        Paragraph("<b>Rcpt</b>", style_tbl_hdr),
        Paragraph("<b>Reg</b>", style_tbl_hdr),
        Paragraph("<b>Member / Contributor Name</b>", style_tbl_hdr),
        Paragraph("<b>Remarks / Dedication</b>", style_tbl_hdr),
        Paragraph("<b>Cash (Rs.)</b>", style_tbl_hdr),
        Paragraph("<b>Bank (Rs.)</b>", style_tbl_hdr),
        Paragraph("<b>Total (Rs.)</b>", style_tbl_hdr),
    ]
]

for idx, x in enumerate(ovbs_list, 1):
    c_str = f"{x['cash']:,.2f}" if x['cash'] > 0 else "-"
    b_str = f"{x['bank']:,.2f}" if x['bank'] > 0 else "-"
    tot_str = f"{x['amt']:,.2f}"
    ovbs_pdf_data.append([
        Paragraph(str(idx), style_tbl_cell_center),
        Paragraph(x['date'], style_tbl_cell_center),
        Paragraph(x['rcpt'], style_tbl_cell_center),
        Paragraph(x['reg'] or "-", style_tbl_cell_center),
        Paragraph(x['name'], style_tbl_cell),
        Paragraph(x['rem'] or "-", style_tbl_cell),
        Paragraph(c_str, style_tbl_cell_right),
        Paragraph(b_str, style_tbl_cell_right),
        Paragraph(f"<b>{tot_str}</b>", style_tbl_cell_bold_right),
    ])

ovbs_pdf_data.append([
    Paragraph("", style_tbl_cell),
    Paragraph("<b>TOTAL</b>", style_tbl_cell_center),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("<b>24 Contributors</b>", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph(f"<b>{tot_ovbs_cash:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{tot_ovbs_bank:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{tot_ovbs_amt:,.2f}</b>", style_tbl_cell_bold_right),
])

ovbs_tbl = Table(ovbs_pdf_data, colWidths=[18, 52, 28, 25, 135, 105, 55, 55, 50], repeatRows=1)
ovbs_tbl_styles = [
    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1B365D')),
    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ('BOTTOMPADDING', (0,0), (-1,-1), 2),
    ('TOPPADDING', (0,0), (-1,-1), 2),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
    ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#E2E8F0')),
]
for r_i in range(1, len(ovbs_list) + 1):
    if r_i % 2 == 0:
        ovbs_tbl_styles.append(('BACKGROUND', (0, r_i), (-1, r_i), colors.HexColor('#F8FAFC')))
ovbs_tbl.setStyle(TableStyle(ovbs_tbl_styles))
story.append(ovbs_tbl)
story.append(Spacer(1, 10))

# 3. CANDLE CUPS & CANDLES TABLE
story.append(Paragraph(f"3. CANDLE CUP & CANDLE SPONSORSHIPS (13 Records | Total: Rs. {candle_sub_amt:,.2f})", style_sec_heading))

candle_pdf_data = [
    [
        Paragraph("<b>#</b>", style_tbl_hdr),
        Paragraph("<b>Date</b>", style_tbl_hdr),
        Paragraph("<b>Rcpt</b>", style_tbl_hdr),
        Paragraph("<b>Reg</b>", style_tbl_hdr),
        Paragraph("<b>Member / Contributor Name</b>", style_tbl_hdr),
        Paragraph("<b>Remarks / Dedication</b>", style_tbl_hdr),
        Paragraph("<b>Cash (Rs.)</b>", style_tbl_hdr),
        Paragraph("<b>Bank (Rs.)</b>", style_tbl_hdr),
        Paragraph("<b>Total (Rs.)</b>", style_tbl_hdr),
    ]
]

# Section A: Candle Cup
candle_pdf_data.append([
    Paragraph("", style_tbl_cell),
    Paragraph("<b>PART A</b>", style_tbl_cell_center),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("<b>CANDLE CUP SPONSORSHIPS (Remarks: 'Candle Cup')</b>", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
])
part_a_header_idx = len(candle_pdf_data) - 1

for idx, x in enumerate(candle_cup_list, 1):
    c_str = f"{x['cash']:,.2f}" if x['cash'] > 0 else "-"
    b_str = f"{x['bank']:,.2f}" if x['bank'] > 0 else "-"
    tot_str = f"{x['amt']:,.2f}"
    candle_pdf_data.append([
        Paragraph(str(idx), style_tbl_cell_center),
        Paragraph(x['date'], style_tbl_cell_center),
        Paragraph(x['rcpt'], style_tbl_cell_center),
        Paragraph(x['reg'] or "-", style_tbl_cell_center),
        Paragraph(x['name'], style_tbl_cell),
        Paragraph(x['rem'] or "-", style_tbl_cell),
        Paragraph(c_str, style_tbl_cell_right),
        Paragraph(b_str, style_tbl_cell_right),
        Paragraph(f"<b>{tot_str}</b>", style_tbl_cell_bold_right),
    ])

candle_pdf_data.append([
    Paragraph("", style_tbl_cell),
    Paragraph("<b>SUBTOTAL</b>", style_tbl_cell_center),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("<b>Subtotal: Candle Cups (7 Contributors)</b>", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph(f"<b>{tot_cc_cash:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{tot_cc_bank:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{tot_cc_amt:,.2f}</b>", style_tbl_cell_bold_right),
])
subtotal_a_idx = len(candle_pdf_data) - 1

# Section B: Other Candles
candle_pdf_data.append([
    Paragraph("", style_tbl_cell),
    Paragraph("<b>PART B</b>", style_tbl_cell_center),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("<b>ADDITIONAL CANDLE SPONSORSHIPS (04-10-2026)</b>", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
])
part_b_header_idx = len(candle_pdf_data) - 1

for idx, x in enumerate(candles_other_list, 1):
    c_str = f"{x['cash']:,.2f}" if x['cash'] > 0 else "-"
    b_str = f"{x['bank']:,.2f}" if x['bank'] > 0 else "-"
    tot_str = f"{x['amt']:,.2f}"
    candle_pdf_data.append([
        Paragraph(str(idx), style_tbl_cell_center),
        Paragraph(x['date'], style_tbl_cell_center),
        Paragraph(x['rcpt'], style_tbl_cell_center),
        Paragraph(x['reg'] or "-", style_tbl_cell_center),
        Paragraph(x['name'], style_tbl_cell),
        Paragraph(x['rem'] or "-", style_tbl_cell),
        Paragraph(c_str, style_tbl_cell_right),
        Paragraph(b_str, style_tbl_cell_right),
        Paragraph(f"<b>{tot_str}</b>", style_tbl_cell_bold_right),
    ])

candle_pdf_data.append([
    Paragraph("", style_tbl_cell),
    Paragraph("<b>SUBTOTAL</b>", style_tbl_cell_center),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("<b>Subtotal: Other Candles (6 Contributors)</b>", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph(f"<b>{tot_oc_cash:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{tot_oc_bank:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{tot_oc_amt:,.2f}</b>", style_tbl_cell_bold_right),
])
subtotal_b_idx = len(candle_pdf_data) - 1

# Grand Combined Candle Total
candle_pdf_data.append([
    Paragraph("", style_tbl_cell),
    Paragraph("<b>TOTAL</b>", style_tbl_cell_center),
    Paragraph("", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph("<b>TOTAL ALL CANDLE CONTRIBUTIONS</b>", style_tbl_cell),
    Paragraph("", style_tbl_cell),
    Paragraph(f"<b>{candle_sub_cash:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{candle_sub_bank:,.2f}</b>", style_tbl_cell_bold_right),
    Paragraph(f"<b>{candle_sub_amt:,.2f}</b>", style_tbl_cell_bold_right),
])
grand_candle_idx = len(candle_pdf_data) - 1

candle_tbl = Table(candle_pdf_data, colWidths=[18, 52, 28, 25, 135, 105, 55, 55, 50], repeatRows=1)
candle_tbl_styles = [
    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1B365D')),
    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ('BOTTOMPADDING', (0,0), (-1,-1), 2),
    ('TOPPADDING', (0,0), (-1,-1), 2),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
    ('BACKGROUND', (0, part_a_header_idx), (-1, part_a_header_idx), colors.HexColor('#E8F1F9')),
    ('BACKGROUND', (0, subtotal_a_idx), (-1, subtotal_a_idx), colors.HexColor('#F1F5F9')),
    ('BACKGROUND', (0, part_b_header_idx), (-1, part_b_header_idx), colors.HexColor('#E8F1F9')),
    ('BACKGROUND', (0, subtotal_b_idx), (-1, subtotal_b_idx), colors.HexColor('#F1F5F9')),
    ('BACKGROUND', (0, grand_candle_idx), (-1, grand_candle_idx), colors.HexColor('#E2E8F0')),
]
candle_tbl.setStyle(TableStyle(candle_tbl_styles))
story.append(candle_tbl)
story.append(Spacer(1, 12))

# Signature block
sig_data = [
    [
        Paragraph("<b>Prepared By</b><br/><br/><br/>Accountant / Operator", style_tbl_cell_center),
        Paragraph("<b>Verified By</b><br/><br/><br/>Trustee / Secretary", style_tbl_cell_center),
        Paragraph("<b>Approved By</b><br/><br/><br/>Vicar / President", style_tbl_cell_center),
    ]
]
sig_tbl = Table(sig_data, colWidths=[174, 175, 174])
sig_tbl.setStyle(TableStyle([
    ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
    ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ('TOPPADDING', (0,0), (-1,-1), 6),
    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
]))
story.append(KeepTogether(sig_tbl))

doc.build(story, canvasmaker=NumberedCanvas)
shutil.copy(PDF_PATH_APP, PDF_PATH_DL)
print(f"PDF report updated successfully.")
